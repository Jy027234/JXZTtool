"""Convert a completed CAAC M1 review workbook into review-truth JSON.

The workbook is only an operator interface.  Candidate values are never used
as expected truth; the ``Truth ...`` columns must be filled from the source
PDF.  Missing values remain missing so the evaluator keeps the hard gates
closed.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any, Mapping


REVIEW_TRUTH_SCHEMA_VERSION = "caac-approved-catalog.v1.review-truth"
REQUIRED_HEADERS = {
    "Review ID",
    "Review type",
    "Section",
    "PDF page",
    "Row",
    "Review status",
    "Reviewer",
    "Notes",
    "Truth row anchor",
    "Truth approval original",
    "Truth holder original",
    "Truth PMA part original",
    "Truth replaced part original",
    "Truth continuation disposition",
}


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _optional_text(value: Any) -> str | None:
    value_text = _text(value)
    return value_text or None


def _optional_bool(value: Any, *, field: str, review_id: str) -> bool | None:
    if value is None or _text(value) == "":
        return None
    if isinstance(value, bool):
        return value
    normalized = _text(value).lower()
    if normalized == "true":
        return True
    if normalized == "false":
        return False
    raise ValueError(f"{review_id}: {field} must be true, false, or blank")


def _review_id(entry: Mapping[str, Any]) -> str:
    explicit = _text(entry.get("reviewId"))
    if explicit:
        return explicit
    row = entry.get("rowNumber")
    return ":".join(
        (
            _text(entry.get("sectionKey")),
            str(int(entry.get("pdfPage") or 0)),
            str(row if row is not None else "page"),
            _text(entry.get("reviewType")),
        )
    )


def _load_json(path: Path) -> dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError(f"expected JSON object: {path}")
    return payload


def _load_sheet_rows(workbook_path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - optional parser dependency
        raise RuntimeError("openpyxl is required to import the review workbook") from exc

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "Review Queue" not in workbook.sheetnames:
        raise ValueError("workbook is missing the Review Queue sheet")
    sheet = workbook["Review Queue"]
    values = list(sheet.iter_rows(values_only=True))
    if len(values) < 5:
        raise ValueError("Review Queue sheet has no row header")
    header_row = [_text(value) for value in values[4]]
    header_index = {name: index for index, name in enumerate(header_row) if name}
    missing_headers = sorted(REQUIRED_HEADERS - set(header_index))
    if missing_headers:
        raise ValueError(f"Review Queue is missing headers: {', '.join(missing_headers)}")

    rows: list[dict[str, Any]] = []
    for raw_row in values[5:]:
        review_id = _text(raw_row[header_index["Review ID"]])
        if not review_id:
            continue
        rows.append(
            {
                header: raw_row[index] if index < len(raw_row) else None
                for header, index in header_index.items()
            }
        )
    return rows


def import_workbook(*, workbook_path: Path, template_path: Path) -> dict[str, Any]:
    template = _load_json(template_path)
    if template.get("schemaVersion") != REVIEW_TRUTH_SCHEMA_VERSION:
        raise ValueError("template schemaVersion is not caac-approved-catalog.v1.review-truth")
    entries = template.get("entries")
    if not isinstance(entries, list):
        raise ValueError("template entries must be an array")

    template_by_id: dict[str, dict[str, Any]] = {}
    for entry in entries:
        if not isinstance(entry, dict):
            raise ValueError("template entries must contain objects")
        key = _review_id(entry)
        if key in template_by_id:
            raise ValueError(f"template has duplicate reviewId: {key}")
        template_by_id[key] = entry

    workbook_rows = _load_sheet_rows(workbook_path)
    seen_ids: set[str] = set()
    for row in workbook_rows:
        review_id = _text(row.get("Review ID"))
        if review_id in seen_ids:
            raise ValueError(f"workbook has duplicate Review ID: {review_id}")
        seen_ids.add(review_id)
        entry = template_by_id.get(review_id)
        if entry is None:
            raise ValueError(f"workbook Review ID is not in template: {review_id}")

        expected = entry.setdefault("expected", {})
        review = entry.setdefault("review", {})
        row_anchor = _optional_bool(row.get("Truth row anchor"), field="Truth row anchor", review_id=review_id)
        if row_anchor is not None:
            expected["rowAnchor"] = row_anchor
        field_mapping = {
            "Truth approval original": "approvalNumberOriginal",
            "Truth holder original": "holderNameOriginal",
            "Truth PMA part original": "pmaPartNumberOriginal",
            "Truth replaced part original": "pmaReplacedPartNumberOriginal",
            "Truth continuation disposition": "continuationDisposition",
        }
        for workbook_field, truth_field in field_mapping.items():
            value = _optional_text(row.get(workbook_field))
            if value is not None:
                expected[truth_field] = value
        review["status"] = _text(row.get("Review status")) or "pending"
        review["reviewer"] = _optional_text(row.get("Reviewer"))
        review["notes"] = _text(row.get("Notes"))

    missing_rows = sorted(set(template_by_id) - seen_ids)
    if missing_rows:
        raise ValueError(f"workbook is missing {len(missing_rows)} template review rows")
    completed = all(
        isinstance(entry.get("review"), dict)
        and entry["review"].get("status") == "confirmed"
        for entry in entries
    )
    template["status"] = "completed" if completed else "in_progress"
    template["reviewWorkbook"] = str(workbook_path)
    return template


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    return parser


def main() -> int:
    args = build_parser().parse_args()
    payload = import_workbook(workbook_path=args.workbook, template_path=args.template)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    confirmed = sum(
        1
        for entry in payload.get("entries") or ()
        if (entry.get("review") or {}).get("status") == "confirmed"
    )
    print(json.dumps({"output": str(args.output), "entries": len(payload.get("entries") or ()), "confirmed": confirmed, "status": payload.get("status")}, ensure_ascii=False))
    return 0


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(main())
